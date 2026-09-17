from __future__ import annotations

import argparse
import json
import math
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Literal

import cv2
import easyocr
import fitz
import numpy as np
from PIL import Image, ImageDraw


TextSource = Literal["auto", "ocr", "text-layer", "both"]
BlockType = Literal["text", "caption", "image", "table"]


@dataclass(frozen=True)
class OcrItem:
    text: str
    confidence: float
    bbox: tuple[float, float, float, float]


@dataclass(frozen=True)
class TextLine:
    text: str
    confidence: float
    bbox: tuple[float, float, float, float]
    source: str
    items: tuple[OcrItem, ...] = ()


@dataclass(frozen=True)
class Caption:
    kind: Literal["figure", "table"]
    identifier: str
    label: str
    line_index: int
    line: TextLine


@dataclass(frozen=True)
class VisualRegion:
    bbox: tuple[int, int, int, int]
    source: str
    score: float
    xref: int | None = None


def ensure_pdf_path(pdf_arg: str | None) -> Path:
    if pdf_arg:
        pdf = Path(pdf_arg)
        if not pdf.exists():
            raise FileNotFoundError(f"PDF not found: {pdf}")
        return pdf

    pdfs = sorted(Path("books").glob("*.pdf"))
    if not pdfs:
        raise FileNotFoundError("No PDF files found in ./books")
    return pdfs[0]


def strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFD", value)
    without_marks = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return without_marks.replace("\u0111", "d").replace("\u0110", "D")


def normalize_for_match(value: str) -> str:
    value = strip_accents(value).lower()
    value = value.replace(",", ".").replace(":", ".").replace(";", ".")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def clean_line_text(value: str) -> str:
    value = value.replace("\u00a0", " ")
    value = re.sub(r"\s+", " ", value)
    return value.strip()


def bbox_to_list(bbox: tuple[float, float, float, float] | tuple[int, int, int, int]) -> list[float]:
    return [round(float(value), 2) for value in bbox]


def bbox_area(bbox: tuple[float, float, float, float] | tuple[int, int, int, int]) -> float:
    x0, y0, x1, y1 = bbox
    return max(0.0, float(x1) - float(x0)) * max(0.0, float(y1) - float(y0))


def bbox_center(bbox: tuple[float, float, float, float] | tuple[int, int, int, int]) -> tuple[float, float]:
    x0, y0, x1, y1 = bbox
    return ((float(x0) + float(x1)) / 2, (float(y0) + float(y1)) / 2)


def bbox_iou(
    left: tuple[float, float, float, float] | tuple[int, int, int, int],
    right: tuple[float, float, float, float] | tuple[int, int, int, int],
) -> float:
    lx0, ly0, lx1, ly1 = left
    rx0, ry0, rx1, ry1 = right
    ix0 = max(float(lx0), float(rx0))
    iy0 = max(float(ly0), float(ry0))
    ix1 = min(float(lx1), float(rx1))
    iy1 = min(float(ly1), float(ry1))
    inter = bbox_area((ix0, iy0, ix1, iy1))
    if inter <= 0:
        return 0.0
    union = bbox_area(left) + bbox_area(right) - inter
    return inter / union if union > 0 else 0.0


def clamp_bbox(
    bbox: tuple[int, int, int, int],
    width: int,
    height: int,
    pad: int = 0,
) -> tuple[int, int, int, int]:
    x0, y0, x1, y1 = bbox
    return (
        max(0, x0 - pad),
        max(0, y0 - pad),
        min(width, x1 + pad),
        min(height, y1 + pad),
    )


def render_page(doc: fitz.Document, page_index: int, scale: float, pages_dir: Path, force: bool) -> Path:
    page_no = page_index + 1
    out_path = pages_dir / f"page_{page_no:03d}.jpg"
    if out_path.exists() and not force:
        return out_path

    page = doc[page_index]
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    pix.save(str(out_path))
    return out_path


def read_page_ocr(reader: easyocr.Reader, image_path: Path) -> list[OcrItem]:
    raw_items = reader.readtext(str(image_path), detail=1, paragraph=False)
    items: list[OcrItem] = []
    for box, text, confidence in raw_items:
        text = clean_line_text(str(text))
        if not text:
            continue
        xs = [float(point[0]) for point in box]
        ys = [float(point[1]) for point in box]
        items.append(
            OcrItem(
                text=text,
                confidence=float(confidence),
                bbox=(min(xs), min(ys), max(xs), max(ys)),
            )
        )
    return items


def group_ocr_lines(items: Iterable[OcrItem], line_tolerance: float) -> list[TextLine]:
    sorted_items = sorted(items, key=lambda item: ((item.bbox[1] + item.bbox[3]) / 2, item.bbox[0]))
    groups: list[list[OcrItem]] = []

    for item in sorted_items:
        center_y = (item.bbox[1] + item.bbox[3]) / 2
        if not groups:
            groups.append([item])
            continue

        current = groups[-1]
        group_center_y = sum((entry.bbox[1] + entry.bbox[3]) / 2 for entry in current) / len(current)
        if abs(center_y - group_center_y) <= line_tolerance:
            current.append(item)
        else:
            groups.append([item])

    lines: list[TextLine] = []
    for group in groups:
        ordered = tuple(sorted(group, key=lambda item: item.bbox[0]))
        x0 = min(item.bbox[0] for item in ordered)
        y0 = min(item.bbox[1] for item in ordered)
        x1 = max(item.bbox[2] for item in ordered)
        y1 = max(item.bbox[3] for item in ordered)
        confidence = sum(item.confidence for item in ordered) / len(ordered)
        text = clean_line_text(" ".join(item.text for item in ordered))
        if text:
            lines.append(TextLine(text=text, confidence=confidence, bbox=(x0, y0, x1, y1), source="easyocr", items=ordered))
    return lines


def extract_text_layer_lines(page: fitz.Page, scale: float) -> list[TextLine]:
    records = page.get_text("dict", sort=True)
    lines: list[TextLine] = []
    for block in records.get("blocks", []):
        if block.get("type") != 0:
            continue
        for raw_line in block.get("lines", []):
            spans = [clean_line_text(span.get("text", "")) for span in raw_line.get("spans", [])]
            text = clean_line_text(" ".join(span for span in spans if span))
            if not text:
                continue
            x0, y0, x1, y1 = raw_line["bbox"]
            lines.append(
                TextLine(
                    text=text,
                    confidence=1.0,
                    bbox=(x0 * scale, y0 * scale, x1 * scale, y1 * scale),
                    source="text-layer",
                )
            )
    return lines


def choose_text_lines(text_source: TextSource, text_lines: list[TextLine], ocr_lines: list[TextLine]) -> tuple[list[TextLine], str]:
    text_chars = sum(len(line.text) for line in text_lines)
    if text_source == "text-layer":
        return text_lines, "text-layer"
    if text_source == "ocr":
        return ocr_lines, "ocr"
    if text_source == "both":
        merged = sorted([*text_lines, *ocr_lines], key=lambda line: (line.bbox[1], line.bbox[0]))
        return merged, "both"
    if text_chars >= 80:
        return text_lines, "text-layer"
    return ocr_lines, "ocr"


def sort_reading_order(lines: list[TextLine], page_width: int) -> list[TextLine]:
    if len(lines) < 12:
        return sorted(lines, key=lambda line: (line.bbox[1], line.bbox[0]))

    centers = [bbox_center(line.bbox)[0] for line in lines]
    left_count = sum(1 for value in centers if value < page_width * 0.47)
    right_count = sum(1 for value in centers if value > page_width * 0.53)
    has_two_columns = left_count >= len(lines) * 0.25 and right_count >= len(lines) * 0.25
    if not has_two_columns:
        return sorted(lines, key=lambda line: (line.bbox[1], line.bbox[0]))

    def key(line: TextLine) -> tuple[int, float, float]:
        center_x, _ = bbox_center(line.bbox)
        column = 0 if center_x < page_width * 0.52 else 1
        return (column, line.bbox[1], line.bbox[0])

    return sorted(lines, key=key)


def is_page_number_noise(line: TextLine, page_height: int) -> bool:
    normalized = normalize_for_match(line.text)
    y0, y1 = line.bbox[1], line.bbox[3]
    near_bottom = y0 > page_height * 0.93
    near_top = y1 < page_height * 0.035
    if not (near_bottom or near_top):
        return False
    if re.fullmatch(r"(trang\s*)?\d{1,4}", normalized):
        return True
    if re.fullmatch(r"\d{1,4}\s*[-_]+\s*", normalized):
        return True
    return False


def caption_from_line(line: TextLine, line_index: int) -> Caption | None:
    normalized = normalize_for_match(line.text)
    non_caption_prefixes = (
        "quan sat",
        "su kien trong",
        "phan biet",
        "cac hinh",
        "tu hinh",
        "trong hinh",
        "duoi day",
        "dua vao",
        "em hay",
    )
    match = re.search(r"\bh[il1]nh\s+(\d{1,2})\s*\.\s*(\d{1,2})\b", normalized)
    if match:
        prefix = normalized[: match.start()].strip(" .:-0123456789")
        if len(prefix) <= 14 and not any(prefix.startswith(value) for value in non_caption_prefixes):
            identifier = f"{int(match.group(1))}_{int(match.group(2))}"
            label = f"Hinh {int(match.group(1))}.{int(match.group(2))}"
            return Caption(kind="figure", identifier=identifier, label=label, line_index=line_index, line=line)

    table_match = re.search(r"\bbang\s+(\d{1,2})(?:\s*\.\s*(\d{1,2}))?\b", normalized)
    if table_match:
        prefix = normalized[: table_match.start()].strip(" .:-0123456789")
        if len(prefix) <= 14 and not any(prefix.startswith(value) for value in non_caption_prefixes):
            minor = table_match.group(2)
            identifier = f"{int(table_match.group(1))}_{int(minor)}" if minor else f"{int(table_match.group(1))}"
            label = f"Bang {identifier.replace('_', '.')}"
            return Caption(kind="table", identifier=identifier, label=label, line_index=line_index, line=line)
    return None


def detect_captions(lines: list[TextLine]) -> list[Caption]:
    captions: list[Caption] = []
    for index, line in enumerate(lines):
        item_captions: list[Caption] = []
        for item in line.items:
            item_line = TextLine(
                text=item.text,
                confidence=item.confidence,
                bbox=item.bbox,
                source=line.source,
                items=(item,),
            )
            caption = caption_from_line(item_line, index)
            if caption and caption.identifier not in {entry.identifier for entry in item_captions}:
                item_captions.append(caption)
        if item_captions:
            captions.extend(item_captions)
            continue

        caption = caption_from_line(line, index)
        if caption:
            captions.append(caption)
    return captions


def detect_visual_regions(
    page_image_path: Path,
    text_lines: list[TextLine],
    min_area_ratio: float,
) -> list[VisualRegion]:
    image_bgr = cv2.imread(str(page_image_path))
    if image_bgr is None:
        return []

    height, width = image_bgr.shape[:2]
    gray = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2GRAY)
    hsv = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2HSV)
    saturation = hsv[:, :, 1]
    edges = cv2.Canny(gray, 55, 160)
    edges = cv2.dilate(edges, np.ones((3, 3), np.uint8), iterations=1)

    color_mask = saturation > 24
    dark_edge_mask = (gray < 218) & (edges > 0)
    mask = np.where(color_mask | dark_edge_mask, 255, 0).astype(np.uint8)

    for line in text_lines:
        x0, y0, x1, y1 = [int(round(value)) for value in line.bbox]
        x0, y0, x1, y1 = clamp_bbox((x0, y0, x1, y1), width, height, pad=6)
        cv2.rectangle(mask, (x0, y0), (x1, y1), 0, thickness=-1)

    mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, np.ones((17, 17), np.uint8), iterations=1)
    mask = cv2.dilate(mask, np.ones((19, 19), np.uint8), iterations=1)

    contours, _ = cv2.findContours(mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    regions: list[VisualRegion] = []
    page_area = width * height
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        area = w * h
        if w < 70 or h < 55:
            continue
        if area < page_area * min_area_ratio:
            continue
        if area > page_area * 0.82:
            continue
        aspect = max(w / max(h, 1), h / max(w, 1))
        if aspect > 13:
            continue
        score = area / page_area
        regions.append(VisualRegion(bbox=(x, y, x + w, y + h), source="cv", score=score))

    return merge_visual_regions(regions, width, height)


def merge_visual_regions(regions: list[VisualRegion], page_width: int, page_height: int) -> list[VisualRegion]:
    merged: list[VisualRegion] = []
    for region in sorted(regions, key=lambda item: bbox_area(item.bbox), reverse=True):
        if any(bbox_iou(region.bbox, existing.bbox) > 0.55 for existing in merged):
            continue
        merged.append(
            VisualRegion(
                bbox=clamp_bbox(region.bbox, page_width, page_height, pad=4),
                source=region.source,
                score=region.score,
                xref=region.xref,
            )
        )
    return sorted(merged, key=lambda item: (item.bbox[1], item.bbox[0]))


def extract_embedded_image_regions(page: fitz.Page, scale: float, page_width: int, page_height: int) -> list[VisualRegion]:
    page_area = page_width * page_height
    regions: list[VisualRegion] = []
    for image_info in page.get_images(full=True):
        xref = int(image_info[0])
        for rect in page.get_image_rects(xref):
            bbox = (
                int(rect.x0 * scale),
                int(rect.y0 * scale),
                int(rect.x1 * scale),
                int(rect.y1 * scale),
            )
            area = bbox_area(bbox)
            if area < page_area * 0.004 or area > page_area * 0.82:
                continue
            regions.append(VisualRegion(bbox=clamp_bbox(bbox, page_width, page_height), source="embedded", score=area / page_area, xref=xref))
    return regions


def fallback_caption_crop(caption: Caption, width: int, height: int) -> tuple[int, int, int, int]:
    cap_x0, cap_y0, cap_x1, cap_y1 = caption.line.bbox
    cap_center_x = (cap_x0 + cap_x1) / 2
    cap_width = cap_x1 - cap_x0
    margin_x = max(24, int(width * 0.04))

    if caption.kind == "table":
        x0 = margin_x
        x1 = width - margin_x
        y0 = int(cap_y1 + height * 0.01)
        y1 = min(height, y0 + int(height * 0.28))
        return clamp_bbox((x0, y0, x1, y1), width, height)

    if cap_width < width * 0.25:
        x0 = max(margin_x, int(cap_center_x - width * 0.24))
        x1 = min(width - margin_x, int(cap_center_x + width * 0.24))
    else:
        x0 = max(margin_x, int(cap_x0 - width * 0.06))
        x1 = min(width - margin_x, int(cap_x1 + width * 0.06))
    y1 = max(1, int(cap_y0 - height * 0.008))
    crop_height = int(height * (0.24 if cap_width < width * 0.27 else 0.32))
    y0 = max(1, y1 - crop_height)
    return clamp_bbox((x0, y0, x1, y1), width, height)


def assign_region_to_caption(
    caption: Caption,
    regions: list[VisualRegion],
    used_region_indexes: set[int],
    page_width: int,
    page_height: int,
) -> tuple[tuple[int, int, int, int], str, int | None]:
    cap_box = caption.line.bbox
    cap_x, _ = bbox_center(cap_box)
    cap_height = max(1.0, cap_box[3] - cap_box[1])
    candidates: list[tuple[float, int, VisualRegion]] = []

    for index, region in enumerate(regions):
        if index in used_region_indexes:
            continue
        rx, _ = bbox_center(region.bbox)
        horizontal = abs(rx - cap_x) / max(page_width, 1)
        if caption.kind == "table":
            vertical_gap = min(
                abs(region.bbox[1] - cap_box[3]),
                abs(cap_box[1] - region.bbox[3]),
            )
            direction_penalty = 0 if region.bbox[1] >= cap_box[1] else 0.12
        else:
            vertical_gap = abs(cap_box[1] - region.bbox[3])
            direction_penalty = 0 if region.bbox[3] <= cap_box[1] + cap_height * 2 else 0.25
        if vertical_gap > page_height * 0.34:
            continue
        if horizontal > 0.33:
            continue
        score = horizontal * 1.7 + vertical_gap / max(page_height, 1) + direction_penalty - region.score * 0.25
        candidates.append((score, index, region))

    if candidates:
        _, index, region = min(candidates, key=lambda item: item[0])
        used_region_indexes.add(index)
        return region.bbox, region.source, index

    return fallback_caption_crop(caption, page_width, page_height), "caption-fallback", None


def save_crop(page_image_path: Path, bbox: tuple[int, int, int, int], output_path: Path) -> None:
    image = Image.open(page_image_path).convert("RGB")
    width, height = image.size
    x0, y0, x1, y1 = clamp_bbox(bbox, width, height, pad=8)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    image.crop((x0, y0, x1, y1)).save(output_path, quality=92)


def image_markdown_alt(block: dict[str, object]) -> str:
    label = str(block.get("label") or block.get("id") or "image")
    page = block.get("page")
    return f"{label} - page {page}"


def block_text_type(line: TextLine) -> BlockType:
    return "caption" if caption_from_line(line, 0) else "text"


def line_markdown_prefix(line: TextLine) -> str | None:
    normalized = normalize_for_match(line.text)
    if re.match(r"^(bai|chuong|phan)\s+\d+", normalized):
        return "###"
    if re.match(r"^\d+\.\d+(\.\d+)?\s+", normalized):
        return "####"
    if re.match(r"^\d+\.\s+\S+", normalized):
        return "####"
    return None


def text_block_to_markdown(block: dict[str, object]) -> list[str]:
    text = clean_line_text(str(block.get("text", "")))
    if not text:
        return []
    if block.get("type") == "caption":
        return [f"*{text}*"]
    prefix = line_markdown_prefix(
        TextLine(
            text=text,
            confidence=float(block.get("confidence", 0.0)),
            bbox=tuple(float(value) for value in block.get("bbox", [0, 0, 0, 0])),  # type: ignore[arg-type]
            source=str(block.get("source", "")),
        )
    )
    if prefix:
        return [f"{prefix} {text}"]
    return [text]


def build_page_blocks(
    lines: list[TextLine],
    captions: list[Caption],
    image_blocks: list[dict[str, object]],
    page_height: int,
) -> list[dict[str, object]]:
    image_by_caption_index: dict[int, list[dict[str, object]]] = {}
    for block in image_blocks:
        line_index = block.get("caption_line_index")
        if isinstance(line_index, int):
            image_by_caption_index.setdefault(line_index, []).append(block)

    skip_line_indexes: set[int] = set()
    image_boxes = [
        tuple(int(value) for value in block["bbox"])  # type: ignore[index]
        for block in image_blocks
        if isinstance(block.get("bbox"), list)
    ]
    for index, line in enumerate(lines):
        if index in image_by_caption_index:
            continue
        center = bbox_center(line.bbox)
        for box in image_boxes:
            if box[0] <= center[0] <= box[2] and box[1] <= center[1] <= box[3]:
                skip_line_indexes.add(index)
                break

    blocks: list[dict[str, object]] = []
    order = 0
    for index, line in enumerate(lines):
        if index in skip_line_indexes:
            continue
        for image_block in image_by_caption_index.get(index, []):
            image_block = dict(image_block)
            image_block["order"] = order
            blocks.append(image_block)
            order += 1

        block_type = block_text_type(line)
        blocks.append(
            {
                "type": block_type,
                "order": order,
                "text": line.text,
                "bbox": bbox_to_list(line.bbox),
                "confidence": round(line.confidence, 4),
                "source": line.source,
            }
        )
        order += 1

    for image_block in image_blocks:
        if image_block.get("caption_line_index") is not None:
            continue
        image_block = dict(image_block)
        image_block["order"] = order
        blocks.append(image_block)
        order += 1

    return blocks


def write_markdown(
    pdf: Path,
    out_dir: Path,
    start_page: int,
    end_page: int,
    page_records: list[dict[str, object]],
) -> None:
    lines: list[str] = [
        f"# {pdf.stem}",
        "",
        f"> Source PDF: `{pdf.as_posix()}`",
        f"> Generated at: `{datetime.now(timezone.utc).isoformat()}`",
        f"> Page range: `{start_page}-{end_page}`",
        "",
    ]

    for page_record in page_records:
        page_no = page_record["page_number"]
        lines.extend([f"## PDF Page {page_no}", ""])
        for block in page_record["text_blocks"]:  # type: ignore[index]
            if block["type"] in {"image", "table"}:
                rel_path = str(block["path"])
                lines.append(f"![{image_markdown_alt(block)}]({rel_path})")
                lines.append("")
                continue
            for md_line in text_block_to_markdown(block):
                lines.append(md_line)
            lines.append("")
        lines.append("</break>")
        lines.append("")

    (out_dir / "book.md").write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def write_docx(out_dir: Path, page_records: list[dict[str, object]], title: str) -> bool:
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception:
        return False

    document = Document()
    document.add_heading(title, level=1)
    for page_record in page_records:
        document.add_heading(f"PDF Page {page_record['page_number']}", level=2)
        for block in page_record["text_blocks"]:  # type: ignore[index]
            block_type = block["type"]
            if block_type in {"image", "table"}:
                image_path = out_dir / str(block["path"])
                if image_path.exists():
                    try:
                        document.add_picture(str(image_path), width=Inches(5.6))
                    except Exception:
                        document.add_paragraph(f"[Image: {block['path']}]")
                continue

            text = clean_line_text(str(block.get("text", "")))
            if not text:
                continue
            prefix = line_markdown_prefix(
                TextLine(
                    text=text,
                    confidence=float(block.get("confidence", 0.0)),
                    bbox=tuple(float(value) for value in block.get("bbox", [0, 0, 0, 0])),  # type: ignore[arg-type]
                    source=str(block.get("source", "")),
                )
            )
            if prefix == "###":
                document.add_heading(text, level=3)
            elif prefix == "####":
                document.add_heading(text, level=4)
            else:
                paragraph = document.add_paragraph(text)
                if block_type == "caption":
                    for run in paragraph.runs:
                        run.italic = True
        document.add_paragraph("</break>")

    document.save(str(out_dir / "book.docx"))
    return True


def chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    text = clean_line_text(text)
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        if end < len(text):
            punct = max(text.rfind(".", start, end), text.rfind("?", start, end), text.rfind("!", start, end))
            if punct > start + chunk_size * 0.55:
                end = punct + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]


def write_rag_chunks(out_dir: Path, page_records: list[dict[str, object]], chunk_size: int, overlap: int) -> None:
    records: list[str] = []
    for page_record in page_records:
        blocks = page_record["text_blocks"]  # type: ignore[index]
        text_parts = [
            str(block.get("text", ""))
            for block in blocks
            if block.get("type") in {"text", "caption"} and str(block.get("text", "")).strip()
        ]
        image_refs = [
            {
                "id": block.get("id"),
                "path": block.get("path"),
                "label": block.get("label"),
                "caption": block.get("caption"),
            }
            for block in blocks
            if block.get("type") in {"image", "table"}
        ]
        for index, chunk in enumerate(chunk_text(" ".join(text_parts), chunk_size, overlap), start=1):
            records.append(
                json.dumps(
                    {
                        "chunk_id": f"page_{int(page_record['page_number']):03d}_{index:02d}",
                        "source_pdf": page_record.get("source_pdf"),
                        "page_number": page_record["page_number"],
                        "text": chunk,
                        "images": image_refs,
                    },
                    ensure_ascii=False,
                )
            )
    (out_dir / "rag_chunks.jsonl").write_text("\n".join(records) + ("\n" if records else ""), encoding="utf-8")


def write_jsonl(path: Path, records: Iterable[dict[str, object]]) -> None:
    path.write_text("\n".join(json.dumps(record, ensure_ascii=False) for record in records) + "\n", encoding="utf-8")


def write_contact_sheet(out_dir: Path, image_records: list[dict[str, object]]) -> Path | None:
    if not image_records:
        return None

    preview_dir = out_dir / "_previews"
    preview_dir.mkdir(parents=True, exist_ok=True)
    thumb_width = 260
    label_height = 42
    padding = 12
    columns = 3
    rows = math.ceil(len(image_records) / columns)
    sheet_width = columns * thumb_width + (columns + 1) * padding
    sheet_height = rows * (thumb_width + label_height) + (rows + 1) * padding
    sheet = Image.new("RGB", (sheet_width, sheet_height), "white")
    draw = ImageDraw.Draw(sheet)

    for index, record in enumerate(image_records):
        image_path = out_dir / str(record["path"])
        if not image_path.exists():
            continue
        with Image.open(image_path).convert("RGB") as image:
            image.thumbnail((thumb_width, thumb_width - label_height), Image.Resampling.LANCZOS)
            col = index % columns
            row = index // columns
            x = padding + col * (thumb_width + padding)
            y = padding + row * (thumb_width + label_height + padding)
            sheet.paste(image, (x + (thumb_width - image.width) // 2, y))
            label = f"{record.get('id')} p.{record.get('page_number') or record.get('page')}"
            draw.text((x + 4, y + thumb_width - 12), label, fill=(20, 20, 20))

    out_path = preview_dir / "images_contact.jpg"
    sheet.save(out_path, quality=88)
    return out_path


def make_image_filename(caption: Caption | None, page_no: int, index: int) -> str:
    if caption is None:
        return f"page_{page_no:03d}_visual_{index:02d}.jpg"
    if caption.kind == "figure":
        return f"{caption.identifier}.jpg"
    return f"table_{caption.identifier}.jpg"


def process_page(
    doc: fitz.Document,
    page_index: int,
    pdf: Path,
    out_dir: Path,
    scale: float,
    reader: easyocr.Reader | None,
    args: argparse.Namespace,
) -> dict[str, object]:
    page = doc[page_index]
    page_no = page_index + 1
    pages_dir = out_dir / "pages"
    images_dir = out_dir / "images"
    metadata_dir = out_dir / "metadata"
    page_metadata_dir = metadata_dir / "pages"

    page_image_path = render_page(doc, page_index, scale, pages_dir, args.force)
    image = Image.open(page_image_path)
    page_width, page_height = image.size

    text_layer_lines = extract_text_layer_lines(page, scale)
    text_layer_chars = sum(len(line.text) for line in text_layer_lines)
    ocr_items: list[OcrItem] = []
    ocr_lines: list[TextLine] = []
    ocr_json_path = metadata_dir / f"ocr_page_{page_no:03d}.json"
    if args.text_source in {"auto", "ocr", "both"} and reader is not None:
        if ocr_json_path.exists() and not args.force_ocr and not args.force:
            cached = json.loads(ocr_json_path.read_text(encoding="utf-8"))
            ocr_items = [
                OcrItem(
                    text=str(item["text"]),
                    confidence=float(item["confidence"]),
                    bbox=tuple(float(value) for value in item["bbox"]),
                )
                for item in cached
            ]
        else:
            ocr_items = read_page_ocr(reader, page_image_path)
            ocr_json_path.write_text(
                json.dumps(
                    [{"text": item.text, "confidence": item.confidence, "bbox": bbox_to_list(item.bbox)} for item in ocr_items],
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
        ocr_lines = group_ocr_lines(ocr_items, line_tolerance=max(9.0, scale * 8.0))

    selected_lines, selected_source = choose_text_lines(args.text_source, text_layer_lines, ocr_lines)
    if args.drop_page_numbers:
        selected_lines = [line for line in selected_lines if not is_page_number_noise(line, page_height)]
    selected_lines = sort_reading_order(selected_lines, page_width)

    captions = detect_captions(selected_lines)
    use_embedded_images = args.embedded_images == "always" or (args.embedded_images == "auto" and text_layer_chars >= 80)
    embedded_regions = extract_embedded_image_regions(page, scale, page_width, page_height) if use_embedded_images else []
    cv_regions = detect_visual_regions(page_image_path, selected_lines, args.min_visual_area_ratio)
    regions = merge_visual_regions([*embedded_regions, *cv_regions], page_width, page_height)

    image_blocks: list[dict[str, object]] = []
    used_region_indexes: set[int] = set()
    for index, caption in enumerate(captions, start=1):
        bbox, region_source, region_index = assign_region_to_caption(caption, regions, used_region_indexes, page_width, page_height)
        file_name = make_image_filename(caption, page_no, index)
        image_path = images_dir / file_name
        save_crop(page_image_path, bbox, image_path)
        rel_path = image_path.relative_to(out_dir).as_posix()
        caption_text = selected_lines[caption.line_index].text if caption.line_index < len(selected_lines) else caption.line.text
        image_blocks.append(
            {
                "type": "table" if caption.kind == "table" else "image",
                "id": caption.identifier,
                "label": caption.label,
                "page": page_no,
                "path": rel_path,
                "caption": caption_text,
                "caption_line_index": caption.line_index,
                "bbox": bbox_to_list(bbox),
                "source": region_source,
                "region_index": region_index,
            }
        )

    if args.include_uncaptioned_visuals:
        next_index = len(image_blocks) + 1
        for region_index, region in enumerate(regions):
            if region_index in used_region_indexes:
                continue
            file_name = make_image_filename(None, page_no, next_index)
            image_path = images_dir / file_name
            save_crop(page_image_path, region.bbox, image_path)
            rel_path = image_path.relative_to(out_dir).as_posix()
            image_blocks.append(
                {
                    "type": "image",
                    "id": f"page_{page_no:03d}_visual_{next_index:02d}",
                    "label": f"Visual {next_index}",
                    "page": page_no,
                    "path": rel_path,
                    "caption": "",
                    "caption_line_index": None,
                    "bbox": bbox_to_list(region.bbox),
                    "source": region.source,
                    "region_index": region_index,
                }
            )
            next_index += 1

    page_blocks = build_page_blocks(selected_lines, captions, image_blocks, page_height)
    page_text = " ".join(
        str(block.get("text", ""))
        for block in page_blocks
        if block.get("type") in {"text", "caption"} and str(block.get("text", "")).strip()
    )

    page_record: dict[str, object] = {
        "source_pdf": pdf.as_posix(),
        "page_number": page_no,
        "page_image_path": page_image_path.relative_to(out_dir).as_posix(),
        "width": page_width,
        "height": page_height,
        "text_source": selected_source,
        "text_layer_characters": text_layer_chars,
        "embedded_images_used": use_embedded_images,
        "ocr_line_count": len(ocr_lines),
        "ocr_mean_confidence": round(sum(line.confidence for line in ocr_lines) / len(ocr_lines), 4) if ocr_lines else None,
        "text": page_text,
        "visual_regions": [
            {"bbox": bbox_to_list(region.bbox), "source": region.source, "score": round(region.score, 5), "xref": region.xref}
            for region in regions
        ],
        "text_blocks": page_blocks,
    }
    page_metadata_dir.mkdir(parents=True, exist_ok=True)
    (page_metadata_dir / f"page_{page_no:03d}.json").write_text(json.dumps(page_record, ensure_ascii=False, indent=2), encoding="utf-8")
    return page_record


def build_summary(
    pdf: Path,
    args: argparse.Namespace,
    page_records: list[dict[str, object]],
    start_page: int,
    end_page: int,
) -> dict[str, object]:
    image_count = sum(
        1
        for page in page_records
        for block in page["text_blocks"]  # type: ignore[index]
        if block.get("type") == "image"
    )
    table_count = sum(
        1
        for page in page_records
        for block in page["text_blocks"]  # type: ignore[index]
        if block.get("type") == "table"
    )
    line_count = sum(
        1
        for page in page_records
        for block in page["text_blocks"]  # type: ignore[index]
        if block.get("type") in {"text", "caption"}
    )
    confidences = [
        float(page["ocr_mean_confidence"])
        for page in page_records
        if page.get("ocr_mean_confidence") is not None
    ]
    return {
        "source_pdf": pdf.as_posix(),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "page_range": {"start": start_page, "end": end_page},
        "pages_processed": len(page_records),
        "scale": args.scale,
        "text_source": args.text_source,
        "ocr_engine": args.ocr_engine,
        "stats": {
            "text_blocks": line_count,
            "images": image_count,
            "tables": table_count,
            "mean_ocr_confidence": round(sum(confidences) / len(confidences), 4) if confidences else None,
        },
        "outputs": {
            "markdown": "book.md",
            "docx": "book.docx",
            "rag_chunks": "rag_chunks.jsonl",
            "page_metadata_dir": "metadata/pages",
            "block_metadata": "metadata/blocks.jsonl",
            "image_metadata": "metadata/images.json",
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract Vietnamese scanned textbooks into Markdown, images, DOCX, and RAG metadata.")
    parser.add_argument("--pdf", default=None, help="Input PDF path. Defaults to the first PDF under ./books.")
    parser.add_argument("--out", default="extracted/class_6", help="Output directory.")
    parser.add_argument("--start-page", type=int, default=1, help="1-based first PDF page.")
    parser.add_argument("--end-page", type=int, default=None, help="1-based last PDF page.")
    parser.add_argument("--scale", type=float, default=2.4, help="PDF render scale used before OCR.")
    parser.add_argument("--langs", nargs="+", default=["vi", "en"], help="EasyOCR language list.")
    parser.add_argument("--ocr-engine", choices=["easyocr", "none"], default="easyocr", help="OCR backend. Use none for text-layer-only extraction.")
    parser.add_argument("--easyocr-gpu", action="store_true", help="Run EasyOCR on GPU. Useful on a dedicated RTX server.")
    parser.add_argument("--text-source", choices=["auto", "ocr", "text-layer", "both"], default="auto", help="Choose text-layer, OCR, or automatic fallback.")
    parser.add_argument("--force", action="store_true", help="Regenerate page images and metadata.")
    parser.add_argument("--force-ocr", action="store_true", help="Ignore cached OCR JSON files.")
    parser.add_argument("--drop-page-numbers", action="store_true", default=True, help="Drop obvious page-number-only header/footer lines.")
    parser.add_argument("--keep-page-numbers", action="store_false", dest="drop_page_numbers", help="Keep page-number-only header/footer lines.")
    parser.add_argument("--include-uncaptioned-visuals", action="store_true", help="Also save visual regions without captions.")
    parser.add_argument("--embedded-images", choices=["auto", "always", "never"], default="auto", help="Use direct embedded-image extraction. Auto enables it only when a text layer exists.")
    parser.add_argument("--min-visual-area-ratio", type=float, default=0.004, help="Minimum detected visual area ratio relative to page area.")
    parser.add_argument("--chunk-size", type=int, default=1800, help="RAG chunk size in characters.")
    parser.add_argument("--chunk-overlap", type=int, default=200, help="RAG chunk overlap in characters.")
    parser.add_argument("--no-docx", action="store_true", help="Skip Word export.")
    parser.add_argument("--no-preview", action="store_true", help="Skip contact sheet preview generation.")
    args = parser.parse_args()

    pdf = ensure_pdf_path(args.pdf)
    out_dir = Path(args.out)
    (out_dir / "pages").mkdir(parents=True, exist_ok=True)
    (out_dir / "images").mkdir(parents=True, exist_ok=True)
    (out_dir / "metadata").mkdir(parents=True, exist_ok=True)

    doc = fitz.open(str(pdf))
    start_index = max(0, args.start_page - 1)
    end_index = doc.page_count if args.end_page is None else min(doc.page_count, args.end_page)
    if start_index >= end_index:
        raise ValueError("Invalid page range.")

    reader: easyocr.Reader | None = None
    if args.ocr_engine == "easyocr" and args.text_source != "text-layer":
        reader = easyocr.Reader(args.langs, gpu=args.easyocr_gpu, verbose=False)

    page_records: list[dict[str, object]] = []
    for page_index in range(start_index, end_index):
        page_no = page_index + 1
        print(f"Process page {page_no}/{end_index}")
        page_records.append(process_page(doc, page_index, pdf, out_dir, args.scale, reader, args))

    write_markdown(pdf, out_dir, start_index + 1, end_index, page_records)
    if not args.no_docx:
        write_docx(out_dir, page_records, pdf.stem)
    write_rag_chunks(out_dir, page_records, args.chunk_size, args.chunk_overlap)

    metadata_dir = out_dir / "metadata"
    block_records: list[dict[str, object]] = []
    image_records: list[dict[str, object]] = []
    for page in page_records:
        for block in page["text_blocks"]:  # type: ignore[index]
            record = dict(block)
            record["page_number"] = page["page_number"]
            block_records.append(record)
            if block.get("type") in {"image", "table"}:
                image_records.append(record)
    write_jsonl(metadata_dir / "blocks.jsonl", block_records)
    (metadata_dir / "images.json").write_text(json.dumps(image_records, ensure_ascii=False, indent=2), encoding="utf-8")
    if not args.no_preview:
        preview_path = write_contact_sheet(out_dir, image_records)
        if preview_path:
            print(f"Preview: {preview_path}")

    summary = build_summary(pdf, args, page_records, start_index + 1, end_index)
    (metadata_dir / "book.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Done: {out_dir / 'book.md'}")
    print(f"Pages: {len(page_records)}")
    print(f"Images: {summary['stats']['images']}, tables: {summary['stats']['tables']}")


if __name__ == "__main__":
    main()
