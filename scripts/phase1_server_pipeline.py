from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import subprocess
import sys
import time
import unicodedata
import zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

os.environ.setdefault("USE_TF", "0")
os.environ.setdefault("TRANSFORMERS_NO_TF", "1")
os.environ.setdefault("USE_FLAX", "0")
os.environ.setdefault("TRANSFORMERS_NO_FLAX", "1")
os.environ.setdefault("TF_CPP_MIN_LOG_LEVEL", "3")
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

ROOT_DIR = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = ROOT_DIR / "scripts"
MARKDOWN_IMAGE_LINK_LINE_RE = re.compile(r"^\s*!\[([^\]]*)\]\(([^)]+)\)\s*$", flags=re.MULTILINE)


def run_cmd(cmd: list[str | Path], *, env: dict[str, str] | None = None, check: bool = True) -> subprocess.CompletedProcess[str]:
    printable = " ".join(str(part) for part in cmd)
    print(f"+ {printable}")
    sys.stdout.flush()
    return subprocess.run([str(part) for part in cmd], env=env, check=check, text=True)


def repo_path(path: str | Path) -> Path:
    path = Path(path)
    return path if path.is_absolute() else ROOT_DIR / path


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_jsonl(path: Path, records: Iterable[dict[str, object]]) -> None:
    rows = [json.dumps(record, ensure_ascii=False) for record in records]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(rows) + ("\n" if rows else ""), encoding="utf-8")


def sanitize_env(extra: dict[str, str] | None = None) -> dict[str, str]:
    env = os.environ.copy()
    env["USE_TF"] = "0"
    env["TRANSFORMERS_NO_TF"] = "1"
    env["USE_FLAX"] = "0"
    env["TRANSFORMERS_NO_FLAX"] = "1"
    env["TF_CPP_MIN_LOG_LEVEL"] = "3"
    env.setdefault("PYTHONUTF8", "1")
    env.setdefault("PYTHONIOENCODING", "utf-8")
    if extra:
        env.update(extra)
    return env


def detect_gpus() -> list[int]:
    try:
        import torch

        return list(range(torch.cuda.device_count()))
    except Exception:
        return []


def strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFD", value)
    without_marks = "".join(ch for ch in normalized if unicodedata.category(ch) != "Mn")
    return without_marks.replace("đ", "d").replace("Đ", "D")


def normalize_text(value: str) -> str:
    value = strip_accents(value).lower()
    value = value.replace(",", ".").replace(":", ".")
    return re.sub(r"\s+", " ", value).strip()


def figure_key(text: str) -> str | None:
    match = re.search(r"h(?:ì|i)nh\s+(\d{1,2})\s*[\.,]\s*(\d{1,2})", text, flags=re.IGNORECASE)
    if not match:
        match = re.search(r"hinh\s+(\d{1,2})\s*[\.,]\s*(\d{1,2})", normalize_text(text))
    if not match:
        return None
    return f"{int(match.group(1))}_{int(match.group(2))}"


def markdown_text_only(text: str) -> str:
    lines: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("![") or stripped == "</break>":
            continue
        stripped = re.sub(r"^#{1,6}\s*", "", stripped)
        if re.fullmatch(r"PDF Page\s+\d+", stripped, flags=re.IGNORECASE):
            continue
        lines.append(stripped)
    cleaned = re.sub(r"\s+", " ", " ".join(lines)).strip()
    return "" if cleaned.lower() in {"none", "null", "nan"} else cleaned


def chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        if end < len(text):
            punct = max(text.rfind(".", start, end), text.rfind("?", start, end), text.rfind("!", start, end))
            if punct > start + chunk_size * 0.55:
                end = punct + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]


def render_one_page(args: tuple[str, int, float, str, bool]) -> dict[str, object]:
    import fitz

    pdf_path, page_index, scale, out_dir, force = args
    doc = fitz.open(pdf_path)
    page_no = page_index + 1
    out_path = Path(out_dir) / f"page_{page_no:03d}.jpg"
    if force or not out_path.exists():
        pix = doc[page_index].get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pix.save(str(out_path))
    return {"page_number": page_no, "path": str(out_path)}


def render_pages(pdf: Path, out_dir: Path, start_page: int, end_page: int | None, scale: float, force: bool) -> list[dict[str, object]]:
    import fitz

    pages_dir = out_dir / "pages"
    pages_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(str(pdf))
    start_index = max(0, start_page - 1)
    end_index = doc.page_count if end_page is None else min(doc.page_count, end_page)
    if start_index >= end_index:
        raise ValueError(f"Invalid page range: {start_page}-{end_page}")

    jobs = [(str(pdf), index, scale, str(pages_dir), force) for index in range(start_index, end_index)]
    worker_count = min(os.cpu_count() or 2, max(1, len(jobs)))
    page_images: list[dict[str, object]] = []
    with ProcessPoolExecutor(max_workers=worker_count) as executor:
        futures = [executor.submit(render_one_page, job) for job in jobs]
        for future in as_completed(futures):
            page_images.append(future.result())
    return sorted(page_images, key=lambda item: int(item["page_number"]))


def copy_source_pdf(pdf: Path, out_dir: Path) -> str:
    source_dir = out_dir / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    dst = source_dir / pdf.name
    if not dst.exists() or dst.stat().st_size != pdf.stat().st_size:
        shutil.copy2(pdf, dst)
    return dst.relative_to(out_dir).as_posix()


def next_image_path(images_dir: Path, stem: str, suffix: str) -> Path:
    suffix = suffix if suffix.startswith(".") else f".{suffix}"
    candidate = images_dir / f"{stem}{suffix}"
    index = 2
    while candidate.exists():
        candidate = images_dir / f"{stem}_{index}{suffix}"
        index += 1
    return candidate


def find_mineru_content_list(raw_dir: Path) -> Path | None:
    files = sorted(raw_dir.rglob("*_content_list.json"))
    if not files:
        files = sorted(raw_dir.rglob("content_list.json"))
    return files[0] if files else None


def load_mineru_visuals(raw_dir: Path, out_dir: Path) -> list[dict[str, object]]:
    content_path = find_mineru_content_list(raw_dir)
    if content_path is None:
        return []

    data = json.loads(content_path.read_text(encoding="utf-8"))
    images_dir = out_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    visuals: list[dict[str, object]] = []

    for index, item in enumerate(data):
        item_type = item.get("type")
        if item_type not in {"image", "table", "chart"}:
            continue

        page_no = int(item.get("page_idx", 0)) + 1
        captions = item.get("image_caption") or item.get("table_caption") or item.get("chart_caption") or []
        caption = " ".join(str(part) for part in captions if str(part).strip()).strip()
        src_rel = item.get("img_path") or item.get("image_path")
        rel_out = ""

        if src_rel:
            src = content_path.parent / str(src_rel)
            if not src.exists():
                matches = list(raw_dir.rglob(Path(str(src_rel)).name))
                src = matches[0] if matches else src
            if src.exists():
                key = figure_key(caption)
                stem = key if key else f"mineru_p{page_no:03d}_{index:04d}"
                dst = next_image_path(images_dir, stem, src.suffix or ".jpg")
                shutil.copy2(src, dst)
                rel_out = dst.relative_to(out_dir).as_posix()

        visual_type = "table" if item_type == "table" else "image"
        visual_id = figure_key(caption) or f"{visual_type}_{index:04d}"
        visuals.append(
            {
                "type": visual_type,
                "id": visual_id,
                "label": caption[:100] if caption else f"{visual_type} {index}",
                "page": page_no,
                "path": rel_out,
                "caption": caption,
                "bbox": item.get("bbox"),
                "source": "mineru_pdf_extract_kit",
                "raw_type": item_type,
            }
        )
    return visuals


def run_mineru_layout(args: argparse.Namespace, pdf: Path, out_dir: Path) -> list[dict[str, object]]:
    if args.layout_engine not in {"auto", "mineru"}:
        return []
    mineru_cmd = shutil.which("mineru")
    local_mineru = ROOT_DIR / "bin" / "mineru"
    if mineru_cmd is None and local_mineru.exists():
        mineru_cmd = str(local_mineru)
    if mineru_cmd is None:
        print("MinerU CLI not found; using CV layout fallback.")
        return []

    raw_dir = out_dir / "_mineru_raw"
    raw_dir.mkdir(parents=True, exist_ok=True)
    cmd: list[str | Path] = [
        mineru_cmd,
        "-p",
        pdf,
        "-o",
        raw_dir,
        "-b",
        "pipeline",
        "-m",
        "ocr",
        "-f",
        "false",
        "-t",
        "true",
    ]
    if args.start_page is not None:
        cmd += ["-s", str(max(0, args.start_page - 1))]
    if args.end_page is not None:
        cmd += ["-e", str(max(0, args.end_page - 1))]

    env = sanitize_env(
        {
            "MINERU_PROCESSING_WINDOW_SIZE": str(args.mineru_window_size),
            "MINERU_API_MAX_CONCURRENT_REQUESTS": "1",
            "MINERU_PDF_RENDER_THREADS": str(min(4, os.cpu_count() or 2)),
        }
    )
    if args.gpu_ids:
        env["CUDA_VISIBLE_DEVICES"] = args.gpu_ids

    result = run_cmd(cmd, env=env, check=False)
    if result.returncode != 0:
        print("MinerU failed; using CV layout fallback.")
        return []

    visuals = load_mineru_visuals(raw_dir, out_dir)
    print(f"MinerU visuals: {len(visuals)}")
    return visuals


def run_cv_layout(args: argparse.Namespace, pdf: Path, out_dir: Path) -> list[dict[str, object]]:
    if args.layout_engine == "mineru":
        return []

    script = SCRIPTS_DIR / "extract_schoolbook.py"
    cmd: list[str | Path] = [
        sys.executable,
        script,
        "--pdf",
        pdf,
        "--out",
        out_dir,
        "--start-page",
        str(args.start_page),
        "--scale",
        str(args.scale),
        "--ocr-engine",
        "easyocr",
        "--text-source",
        "ocr",
        "--embedded-images",
        "never",
        "--chunk-size",
        str(args.chunk_size),
        "--chunk-overlap",
        str(args.chunk_overlap),
    ]
    if args.end_page is not None:
        cmd += ["--end-page", str(args.end_page)]
    if args.easyocr_gpu:
        cmd.append("--easyocr-gpu")
    if args.include_uncaptioned_visuals:
        cmd.append("--include-uncaptioned-visuals")
    if args.force:
        cmd.append("--force")
    if args.force_ocr:
        cmd.append("--force-ocr")

    run_cmd(cmd, env=sanitize_env(), check=True)
    images_path = out_dir / "metadata" / "images.json"
    if not images_path.exists():
        return []
    visuals = json.loads(images_path.read_text(encoding="utf-8"))
    for visual in visuals:
        visual.setdefault("page", visual.get("page_number"))
        visual["source"] = visual.get("source") or "cv_caption_fallback"
    print(f"CV visuals: {len(visuals)}")
    return visuals


def render_manifest(page_images: list[dict[str, object]], manifest_path: Path) -> None:
    records = []
    for page in page_images:
        records.append({"page_number": int(page["page_number"]), "path": str(Path(str(page["path"])).resolve())})
    write_jsonl(manifest_path, records)


def split_records(records: list[dict[str, object]], shard_count: int) -> list[list[dict[str, object]]]:
    shards: list[list[dict[str, object]]] = [[] for _ in range(max(1, shard_count))]
    for index, record in enumerate(records):
        shards[index % len(shards)].append(record)
    return shards


def run_deepseek_ocr(args: argparse.Namespace, page_images: list[dict[str, object]], out_dir: Path) -> bool:
    if args.skip_deepseek:
        return False

    gpu_ids = [int(part) for part in args.gpu_ids.split(",") if part.strip()] if args.gpu_ids else detect_gpus()
    if not gpu_ids and not args.allow_cpu_deepseek:
        print("No GPU detected; DeepSeek-OCR skipped.")
        return False

    deepseek_dir = out_dir / "_deepseek_pages"
    manifest_dir = out_dir / "_manifests"
    deepseek_dir.mkdir(parents=True, exist_ok=True)
    manifest_dir.mkdir(parents=True, exist_ok=True)

    worker = SCRIPTS_DIR / "deepseek_ocr_pages.py"
    records = [{"page_number": int(page["page_number"]), "path": str(Path(str(page["path"])).resolve())} for page in page_images]
    shards = split_records(records, len(gpu_ids) if gpu_ids else 1)

    processes: list[subprocess.Popen[bytes]] = []
    for shard_index, shard in enumerate(shards):
        if not shard:
            continue
        gpu_id = gpu_ids[shard_index] if gpu_ids else None
        manifest = manifest_dir / f"deepseek_shard_{shard_index}.jsonl"
        write_jsonl(manifest, shard)
        env = sanitize_env(
            {
                "DEEPSEEK_MODEL": args.deepseek_model,
                "DEEPSEEK_PROMPT": args.deepseek_prompt,
            }
        )
        if gpu_id is not None:
            env["CUDA_VISIBLE_DEVICES"] = str(gpu_id)
        cmd: list[str | Path] = [
            sys.executable,
            worker,
            "--manifest",
            manifest,
            "--out",
            deepseek_dir,
            "--model",
            args.deepseek_model,
            "--prompt",
            args.deepseek_prompt,
            "--attn-impl",
            args.deepseek_attn_impl,
            "--base-size",
            str(args.deepseek_base_size),
            "--image-size",
            str(args.deepseek_image_size),
        ]
        if args.force_deepseek or args.force:
            cmd.append("--force")
        print(f"DeepSeek shard {shard_index}: GPU={gpu_id}, pages={len(shard)}")
        processes.append(subprocess.Popen([str(part) for part in cmd], env=env))

    ok = True
    for process in processes:
        ok = (process.wait() == 0) and ok
    return ok


def load_deepseek_pages(out_dir: Path) -> dict[int, dict[str, object]]:
    pages: dict[int, dict[str, object]] = {}
    for path in sorted((out_dir / "_deepseek_pages").glob("page_*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        pages[int(payload["page_number"])] = payload
    return pages


def visual_page(visual: dict[str, object]) -> int:
    return int(visual.get("page") or visual.get("page_number") or 0)


def path_relative_to(path: Path, base: Path) -> str:
    try:
        return path.relative_to(base).as_posix()
    except ValueError:
        return str(path)


def normalize_deepseek_markdown_images(
    page_no: int,
    markdown: str,
    deepseek_payload: dict[str, object],
    out_dir: Path,
) -> str:
    raw_dir_value = str(deepseek_payload.get("raw_output_dir") or "")
    raw_dir = Path(raw_dir_value) if raw_dir_value else None
    images_dir = out_dir / "images"
    images_dir.mkdir(parents=True, exist_ok=True)
    image_index = 0

    def replace_link(match: re.Match[str]) -> str:
        nonlocal image_index
        alt = match.group(1)
        rel_path = match.group(2).strip()
        if rel_path.startswith(("http://", "https://", "data:", "#")):
            return match.group(0)
        if raw_dir is None:
            return ""

        src = raw_dir / rel_path
        if not src.exists():
            matches = list(raw_dir.rglob(Path(rel_path).name))
            src = matches[0] if matches else src
        if not src.exists():
            return ""

        suffix = src.suffix or ".jpg"
        dst = next_image_path(images_dir, f"page_{page_no:03d}_deepseek_{image_index:03d}", suffix)
        shutil.copy2(src, dst)
        image_index += 1
        return f"![{alt}]({dst.relative_to(out_dir).as_posix()})"

    normalized = MARKDOWN_IMAGE_LINK_LINE_RE.sub(replace_link, markdown)
    return re.sub(r"(\]\([^)]+\))(?=!\[)", r"\1\n", normalized)


def image_markdown(visual: dict[str, object]) -> str:
    alt = str(visual.get("caption") or visual.get("label") or visual.get("id") or "image").replace("\n", " ")
    return f"![{alt}]({visual.get('path', '')})"


def insert_images_into_markdown(page_no: int, markdown: str, visuals: list[dict[str, object]]) -> tuple[list[str], list[dict[str, object]]]:
    page_visuals = [dict(visual) for visual in visuals if visual_page(visual) == page_no and visual.get("path")]
    for visual in page_visuals:
        visual["figure_key"] = figure_key(str(visual.get("caption") or visual.get("label") or ""))
        visual["used"] = False

    output_lines: list[str] = []
    blocks: list[dict[str, object]] = []
    order = 0
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        norm_line = normalize_text(line)
        line_key = figure_key(line)
        image_match = MARKDOWN_IMAGE_LINK_LINE_RE.match(line)

        if image_match:
            output_lines.append(line)
            image_path = image_match.group(2).strip()
            image_id = Path(image_path).stem
            blocks.append(
                {
                    "type": "image",
                    "id": image_id,
                    "label": image_match.group(1) or image_id,
                    "caption": image_match.group(1) or "",
                    "order": order,
                    "page": page_no,
                    "path": image_path,
                    "source": "deepseek-ocr-markdown",
                }
            )
            order += 1
            continue

        for visual in page_visuals:
            if visual.get("used"):
                continue
            caption_norm = normalize_text(str(visual.get("caption") or ""))
            hit = bool(visual.get("figure_key") and visual.get("figure_key") == line_key)
            hit = hit or bool(caption_norm and len(caption_norm) > 10 and caption_norm[:45] in norm_line)
            if hit:
                clean_visual = {key: value for key, value in visual.items() if key not in {"used", "figure_key"}}
                clean_visual["order"] = order
                output_lines.append(image_markdown(clean_visual))
                blocks.append(clean_visual)
                visual["used"] = True
                order += 1

        if line.strip():
            output_lines.append(line)
            block_type = "caption" if figure_key(line) else "text"
            blocks.append({"type": block_type, "order": order, "text": line, "page": page_no, "source": "deepseek-ocr"})
            order += 1

    for visual in page_visuals:
        if visual.get("used"):
            continue
        clean_visual = {key: value for key, value in visual.items() if key not in {"used", "figure_key"}}
        clean_visual["order"] = order
        output_lines.append(image_markdown(clean_visual))
        blocks.append(clean_visual)
        order += 1

    return output_lines, blocks


def write_merged_outputs(
    args: argparse.Namespace,
    pdf: Path,
    out_dir: Path,
    page_images: list[dict[str, object]],
    visuals: list[dict[str, object]],
    deepseek_pages: dict[int, dict[str, object]],
    source_pdf_rel: str | None,
) -> list[dict[str, object]]:
    metadata_dir = out_dir / "metadata"
    pages_metadata_dir = metadata_dir / "pages"
    pages_metadata_dir.mkdir(parents=True, exist_ok=True)

    book_lines = [
        f"# {pdf.stem}",
        "",
        f"> Source PDF: `{source_pdf_rel or pdf.as_posix()}`",
        f"> Generated at: `{datetime.now(timezone.utc).isoformat()}`",
        f"> OCR engine: `deepseek-ocr`",
        f"> Layout engine: `{args.layout_engine}`",
        "",
    ]
    page_records: list[dict[str, object]] = []
    block_records: list[dict[str, object]] = []
    rag_records: list[dict[str, object]] = []

    for page in page_images:
        page_no = int(page["page_number"])
        deepseek_payload = deepseek_pages.get(page_no, {})
        page_md = str(deepseek_payload.get("markdown") or "").strip()
        page_md = normalize_deepseek_markdown_images(page_no, page_md, deepseek_payload, out_dir)
        if not page_md:
            page_md = f"[OCR missing for PDF page {page_no}]"

        page_lines, blocks = insert_images_into_markdown(page_no, page_md, visuals)
        book_lines.extend([f"## PDF Page {page_no}", "", *page_lines, "", "</break>", ""])

        page_text = markdown_text_only("\n".join(page_lines))
        rel_page_path = Path(str(page["path"]))
        try:
            rel_page = rel_page_path.relative_to(out_dir).as_posix()
        except ValueError:
            rel_page = str(rel_page_path)

        page_record = {
            "source_pdf": source_pdf_rel or pdf.as_posix(),
            "page_number": page_no,
            "page_image_path": rel_page,
            "text": page_text,
            "text_blocks": blocks,
            "ocr_engine": "deepseek-ocr",
            "ocr_error": deepseek_payload.get("error"),
        }
        page_records.append(page_record)
        write_json(pages_metadata_dir / f"page_{page_no:03d}.json", page_record)

        for block in blocks:
            block_record = dict(block)
            block_record["page_number"] = page_no
            block_records.append(block_record)

        images_for_page = [
            {key: block.get(key) for key in ("id", "path", "caption", "label", "type")}
            for block in blocks
            if block.get("type") in {"image", "table"}
        ]
        for index, chunk in enumerate(chunk_text(page_text, args.chunk_size, args.chunk_overlap), start=1):
            rag_records.append(
                {
                    "chunk_id": f"page_{page_no:03d}_{index:02d}",
                    "source_pdf": source_pdf_rel or pdf.as_posix(),
                    "page_number": page_no,
                    "text": chunk,
                    "images": images_for_page,
                }
            )

    (out_dir / "book.md").write_text("\n".join(book_lines).strip() + "\n", encoding="utf-8")
    write_jsonl(metadata_dir / "blocks.jsonl", block_records)
    image_records = [record for record in block_records if record.get("type") in {"image", "table"}]
    write_json(metadata_dir / "images.json", image_records)
    write_jsonl(out_dir / "rag_chunks.jsonl", rag_records)

    summary = {
        "source_pdf": source_pdf_rel or pdf.as_posix(),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "page_range": {"start": args.start_page, "end": args.end_page},
        "pages_processed": len(page_records),
        "scale": args.scale,
        "ocr_engine": "deepseek-ocr",
        "layout_engine": args.layout_engine,
        "stats": {
            "blocks": len(block_records),
            "images": len(image_records),
            "rag_chunks": len(rag_records),
        },
        "outputs": {
            "markdown": "book.md",
            "docx": "book.docx",
            "rag_chunks": "rag_chunks.jsonl",
            "page_metadata_dir": "metadata/pages",
            "block_metadata": "metadata/blocks.jsonl",
            "image_metadata": "metadata/images.json",
        },
    }
    write_json(metadata_dir / "book.json", summary)
    return page_records


def write_docx_from_markdown(pdf: Path, out_dir: Path) -> bool:
    try:
        from docx import Document
        from docx.shared import Inches
    except Exception:
        print("python-docx is not available; skip DOCX.")
        return False

    book_md = out_dir / "book.md"
    if not book_md.exists():
        return False

    document = Document()
    document.add_heading(pdf.stem, level=1)
    for line in book_md.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped == "</break>" or stripped.startswith("> "):
            continue
        if stripped.startswith("# "):
            continue
        if stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("#### "):
            document.add_heading(stripped[5:].strip(), level=4)
        elif stripped.startswith("!["):
            match = re.search(r"\]\((.*?)\)", stripped)
            if match:
                image_path = out_dir / match.group(1)
                if image_path.exists():
                    try:
                        document.add_picture(str(image_path), width=Inches(5.6))
                    except Exception:
                        document.add_paragraph(str(image_path))
        else:
            document.add_paragraph(stripped)
    document.save(str(out_dir / "book.docx"))
    return True


def write_contact_sheet(out_dir: Path) -> Path | None:
    try:
        from PIL import Image, ImageDraw
    except Exception:
        return None

    images_path = out_dir / "metadata" / "images.json"
    if not images_path.exists():
        return None
    records = [record for record in json.loads(images_path.read_text(encoding="utf-8")) if record.get("path")]
    if not records:
        return None

    preview_dir = out_dir / "_previews"
    preview_dir.mkdir(parents=True, exist_ok=True)
    thumb_width, label_height, padding, columns = 260, 42, 12, 3
    max_rows_per_sheet = 80
    max_records_per_sheet = columns * max_rows_per_sheet
    first_path: Path | None = None

    for sheet_index, start in enumerate(range(0, len(records), max_records_per_sheet), start=1):
        batch = records[start : start + max_records_per_sheet]
        rows = math.ceil(len(batch) / columns)
        sheet = Image.new(
            "RGB",
            (
                columns * thumb_width + (columns + 1) * padding,
                rows * (thumb_width + label_height) + (rows + 1) * padding,
            ),
            "white",
        )
        draw = ImageDraw.Draw(sheet)

        for index, record in enumerate(batch):
            image_path = out_dir / str(record["path"])
            if not image_path.exists():
                continue
            with Image.open(image_path).convert("RGB") as image:
                image.thumbnail((thumb_width, thumb_width - label_height), Image.Resampling.LANCZOS)
                col = index % columns
                row = index // columns
                x = padding + col * (thumb_width + padding)
                y = padding + row * (thumb_width + label_height + padding)
                sheet.paste(image, (x + (thumb_width - image.width) // 2, y))
                label = f"{record.get('id')} p.{record.get('page_number') or record.get('page')}"
                draw.text((x + 4, y + thumb_width - 12), label, fill=(20, 20, 20))

        out_path = preview_dir / ("images_contact.jpg" if sheet_index == 1 else f"images_contact_{sheet_index:03d}.jpg")
        sheet.save(out_path, quality=88)
        first_path = first_path or out_path

    return first_path


def run_spellcheck(args: argparse.Namespace, out_dir: Path) -> bool:
    if args.skip_spellcheck:
        return False
    script = SCRIPTS_DIR / "spellcheck_markdown_qwen.py"
    if not script.exists():
        print("Spellcheck script not found; skip.")
        return False

    corrected_md = out_dir / "book_corrected.md"
    cmd: list[str | Path] = [
        sys.executable,
        script,
        "--input",
        out_dir / "book.md",
        "--output",
        corrected_md,
        "--raw-output",
        out_dir / "book_raw_ocr.md",
        "--metadata-dir",
        out_dir / "metadata",
        "--rag-output",
        out_dir / "rag_chunks.jsonl",
        "--model",
        args.spellcheck_model,
        "--gpus",
        args.gpu_ids or "auto",
        "--max-section-chars",
        str(args.spellcheck_max_section_chars),
        "--max-new-tokens",
        str(args.spellcheck_max_new_tokens),
        "--chunk-size",
        str(args.chunk_size),
        "--chunk-overlap",
        str(args.chunk_overlap),
    ]
    if args.force_spellcheck or args.force:
        cmd.append("--force")

    result = run_cmd(cmd, env=sanitize_env(), check=False)
    if result.returncode == 0 and corrected_md.exists():
        shutil.copy2(corrected_md, out_dir / "book.md")
        return True
    print("Spellcheck failed or produced no output; keep raw OCR markdown.")
    return False


def zip_output(out_dir: Path, zip_path: Path | None) -> Path:
    target = zip_path or out_dir.with_suffix(".zip")
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        target.unlink()
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in sorted(out_dir.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(out_dir.parent))
    return target


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Server RTX 4070 phase-1 textbook extraction pipeline.")
    parser.add_argument("--pdf", type=Path, default=Path("books/SGK Lịch sử và địa lí 6 CD.pdf"))
    parser.add_argument("--out", type=Path, default=Path("extracted/class_6_rtx4070"))
    parser.add_argument("--start-page", type=int, default=1)
    parser.add_argument("--end-page", type=int, default=None)
    parser.add_argument("--scale", type=float, default=2.4)
    parser.add_argument("--layout-engine", choices=["auto", "mineru", "cv", "none"], default="auto")
    parser.add_argument("--easyocr-gpu", action="store_true")
    parser.add_argument("--include-uncaptioned-visuals", action="store_true")
    parser.add_argument("--skip-deepseek", action="store_true")
    parser.add_argument("--allow-cpu-deepseek", action="store_true")
    parser.add_argument("--deepseek-model", default="deepseek-ai/DeepSeek-OCR")
    parser.add_argument("--deepseek-prompt", default="<image>\n<|grounding|>Convert the document to markdown.")
    parser.add_argument("--deepseek-attn-impl", default="sdpa")
    parser.add_argument("--deepseek-base-size", type=int, default=1024)
    parser.add_argument("--deepseek-image-size", type=int, default=640)
    parser.add_argument("--skip-spellcheck", action="store_true")
    parser.add_argument("--spellcheck-model", default="Qwen/Qwen3-4B-Instruct-2507")
    parser.add_argument("--spellcheck-max-section-chars", type=int, default=4200)
    parser.add_argument("--spellcheck-max-new-tokens", type=int, default=3072)
    parser.add_argument("--gpu-ids", default="", help="Comma-separated GPU ids. Empty = auto-detect.")
    parser.add_argument("--mineru-window-size", type=int, default=16)
    parser.add_argument("--chunk-size", type=int, default=1800)
    parser.add_argument("--chunk-overlap", type=int, default=200)
    parser.add_argument("--copy-source-pdf", action="store_true", default=True)
    parser.add_argument("--no-copy-source-pdf", action="store_false", dest="copy_source_pdf")
    parser.add_argument("--zip-output", action="store_true", default=True)
    parser.add_argument("--no-zip-output", action="store_false", dest="zip_output")
    parser.add_argument("--zip-path", type=Path, default=None)
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--force-ocr", action="store_true")
    parser.add_argument("--force-deepseek", action="store_true")
    parser.add_argument("--force-spellcheck", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    pdf = repo_path(args.pdf)
    out_dir = repo_path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "images").mkdir(parents=True, exist_ok=True)
    (out_dir / "metadata").mkdir(parents=True, exist_ok=True)

    if not pdf.exists():
        raise FileNotFoundError(pdf)

    start_time = time.time()
    gpu_ids = [int(part) for part in args.gpu_ids.split(",") if part.strip()] if args.gpu_ids else detect_gpus()
    print(f"PDF: {pdf}")
    print(f"OUT: {out_dir}")
    print(f"GPUs: {gpu_ids}")

    source_pdf_rel = copy_source_pdf(pdf, out_dir) if args.copy_source_pdf else None
    page_images = render_pages(pdf, out_dir, args.start_page, args.end_page, args.scale, args.force)
    render_manifest(page_images, out_dir / "_manifests" / "pages.jsonl")
    print(f"Rendered pages: {len(page_images)}")

    visuals = run_mineru_layout(args, pdf, out_dir)
    if not visuals and args.layout_engine in {"auto", "cv"}:
        visuals = run_cv_layout(args, pdf, out_dir)
    write_json(out_dir / "metadata" / "images.json", visuals)

    deepseek_ok = run_deepseek_ocr(args, page_images, out_dir)
    deepseek_pages = load_deepseek_pages(out_dir)
    if deepseek_pages:
        write_merged_outputs(args, pdf, out_dir, page_images, visuals, deepseek_pages, source_pdf_rel)
    elif (out_dir / "book.md").exists():
        print("DeepSeek output missing; keeping CV/EasyOCR markdown from layout fallback.")
    else:
        raise RuntimeError("DeepSeek output missing and no fallback markdown exists.")

    spellcheck_ok = run_spellcheck(args, out_dir)
    write_docx_from_markdown(pdf, out_dir)
    contact = write_contact_sheet(out_dir)

    book_json_path = out_dir / "metadata" / "book.json"
    if book_json_path.exists():
        summary = json.loads(book_json_path.read_text(encoding="utf-8"))
    else:
        summary = {}
    summary["server_pipeline"] = {
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "seconds": round(time.time() - start_time, 2),
        "gpu_ids": gpu_ids,
        "deepseek_ok": bool(deepseek_ok),
        "spellcheck_ok": bool(spellcheck_ok),
        "contact_sheet": contact.relative_to(out_dir).as_posix() if contact else None,
    }
    write_json(book_json_path, summary)

    zip_path = zip_output(out_dir, repo_path(args.zip_path) if args.zip_path else None) if args.zip_output else None
    print("Done.")
    print(f"book.md: {out_dir / 'book.md'}")
    print(f"book.docx: {out_dir / 'book.docx'}")
    print(f"rag_chunks.jsonl: {out_dir / 'rag_chunks.jsonl'}")
    if contact:
        print(f"contact sheet: {contact}")
    if zip_path:
        print(f"zip: {zip_path}")


if __name__ == "__main__":
    main()
