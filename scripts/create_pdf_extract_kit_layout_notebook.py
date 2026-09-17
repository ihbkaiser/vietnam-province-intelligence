from __future__ import annotations

import json
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "notebooks" / "phase1_kaggle_pdf_extract_kit_layout.ipynb"


def md(source: str) -> dict[str, object]:
    return {"cell_type": "markdown", "metadata": {}, "source": source.splitlines(keepends=True)}


def code(source: str) -> dict[str, object]:
    return {
        "cell_type": "code",
        "execution_count": None,
        "metadata": {},
        "outputs": [],
        "source": source.splitlines(keepends=True),
    }


cells = [
    md(
        """# Phase 1 Kaggle - PDF-Extract-Kit Direct Layout/Image Extraction

Notebook này thử trực tiếp PDF-Extract-Kit, không qua MinerU wrapper. Mục tiêu của lượt này là layout detection và crop ảnh/bảng/caption bằng DocLayout-YOLO từ PDF-Extract-Kit. Text OCR tiếng Việt nên chạy bằng notebook DeepSeek-OCR riêng rồi merge sau."""
    ),
    code(
        r"""from pathlib import Path
import json, math, os, re, subprocess, sys, zipfile
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime, timezone

os.environ.setdefault('USE_TF', '0')
os.environ.setdefault('TRANSFORMERS_NO_TF', '1')
os.environ.setdefault('USE_FLAX', '0')
os.environ.setdefault('TRANSFORMERS_NO_FLAX', '1')
os.environ.setdefault('TF_CPP_MIN_LOG_LEVEL', '3')
os.environ.setdefault('PYTHONUTF8', '1')
os.environ.setdefault('PYTHONIOENCODING', 'utf-8')

# Thay đường dẫn input của file PDF vào đây.
INPUT_PDF = Path('/kaggle/input/vietnam-schoolbooks/SGK Lịch sử và địa lí 6 CD.pdf')
if not INPUT_PDF.exists() and Path('/kaggle/input').exists():
    candidates = sorted(Path('/kaggle/input').rglob('*.pdf'))
    if candidates:
        INPUT_PDF = candidates[0]
if not INPUT_PDF.exists():
    INPUT_PDF = Path('books/SGK Lịch sử và địa lí 6 CD.pdf')

OUTPUT_DIR = Path('/kaggle/working/class_6_pdf_extract_kit_layout') if Path('/kaggle/working').exists() else Path('extracted/class_6_pdf_extract_kit_layout')
KIT_DIR = OUTPUT_DIR / '_PDF-Extract-Kit'
MODEL_DIR = OUTPUT_DIR / '_models_pdf_extract_kit'
PAGES_DIR = OUTPUT_DIR / 'pages'
IMAGES_DIR = OUTPUT_DIR / 'images'
CAPTIONS_DIR = OUTPUT_DIR / 'caption_crops'
METADATA_DIR = OUTPUT_DIR / 'metadata'
PREVIEW_DIR = OUTPUT_DIR / '_previews'
LAYOUT_VIS_DIR = PREVIEW_DIR / 'layout_boxes'

# Test trước pages 6-10. Khi ổn thì đổi START_PAGE=1, END_PAGE=None để chạy toàn bộ sách.
START_PAGE = 6
END_PAGE = 10
RENDER_SCALE = 2.0

LAYOUT_MODEL_REPO_CANDIDATES = ['opendatalab/pdf-extract-kit-1.0', 'opendatalab/PDF-Extract-Kit-1.0']
LAYOUT_WEIGHT_PATTERN = 'models/Layout/YOLO/doclayout_yolo_ft.pt'
IMG_SIZE = 1024
CONF_THRES = 0.25
IOU_THRES = 0.45
DEVICE = 'auto'  # 'auto', '0', 'cpu'

VISUAL_TYPES = {'figure', 'table'}
CAPTION_TYPES = {'figure_caption', 'table_caption'}
SAVE_CAPTION_CROPS = True
INCLUDE_CAPTION_CROPS_IN_MD = False
MIN_CROP_AREA_RATIO = 0.002
PAD_PX = 8
ZIP_OUTPUT = True
CHUNK_SIZE = 1800
CHUNK_OVERLAP = 200

for directory in [OUTPUT_DIR, MODEL_DIR, PAGES_DIR, IMAGES_DIR, CAPTIONS_DIR, METADATA_DIR, PREVIEW_DIR, LAYOUT_VIS_DIR]:
    directory.mkdir(parents=True, exist_ok=True)

print('INPUT_PDF =', INPUT_PDF)
print('OUTPUT_DIR =', OUTPUT_DIR)
print('Page range =', START_PAGE, END_PAGE)
"""
    ),
    code(
        r"""def run_cmd(cmd, *, env=None, check=True):
    print('+', ' '.join(map(str, cmd)))
    sys.stdout.flush()
    return subprocess.run(list(map(str, cmd)), env=env, check=check)

# Cài layout-only, tránh requirements.txt đầy đủ vì nó kéo PaddleOCR/lmdeploy và CUDA stack nặng.
run_cmd([sys.executable, '-m', 'pip', 'install', '-q', '-U', 'pip'])
run_cmd([
    sys.executable, '-m', 'pip', 'install', '-q', '--no-cache-dir',
    'PyMuPDF', 'Pillow', 'opencv-python-headless', 'numpy<2.5',
    'huggingface_hub', 'ultralytics>=8.2.85',
    'omegaconf', 'pyyaml', 'python-docx', 'tqdm'
])
doclayout_install = None
for package in ['doclayout-yolo==0.0.4', 'doclayout-yolo==0.0.3', 'doclayout-yolo==0.0.2b1']:
    doclayout_install = run_cmd([sys.executable, '-m', 'pip', 'install', '-q', '--no-cache-dir', package], check=False)
    if doclayout_install.returncode == 0:
        print('doclayout-yolo installed:', package)
        break
if doclayout_install is None or doclayout_install.returncode != 0:
    print('doclayout-yolo install failed; notebook will fallback to ultralytics.YOLO.')

try:
    import torch
    GPU_COUNT = torch.cuda.device_count()
    GPU_IDS = list(range(GPU_COUNT))
except Exception as exc:
    print('torch import failed:', repr(exc))
    GPU_COUNT = 0
    GPU_IDS = []

CPU_COUNT = os.cpu_count() or 2
os.environ['OMP_NUM_THREADS'] = str(CPU_COUNT)
os.environ['MKL_NUM_THREADS'] = str(CPU_COUNT)
print('CPU_COUNT =', CPU_COUNT)
print('GPU_IDS =', GPU_IDS)
run_cmd(['nvidia-smi'], check=False)
"""
    ),
    code(
        r"""# Clone PDF-Extract-Kit để dùng trực tiếp model/code layout, không dùng MinerU wrapper.
if not KIT_DIR.exists():
    run_cmd(['git', 'clone', '--depth', '1', 'https://github.com/opendatalab/PDF-Extract-Kit.git', str(KIT_DIR)], check=True)
else:
    print('PDF-Extract-Kit repo exists:', KIT_DIR)

if str(KIT_DIR) not in sys.path:
    sys.path.insert(0, str(KIT_DIR))

from huggingface_hub import snapshot_download

weight_path = MODEL_DIR / LAYOUT_WEIGHT_PATTERN
if not weight_path.exists():
    last_error = None
    for repo_id in LAYOUT_MODEL_REPO_CANDIDATES:
        try:
            print('download layout weights from', repo_id)
            snapshot_download(repo_id=repo_id, local_dir=str(MODEL_DIR), allow_patterns=[LAYOUT_WEIGHT_PATTERN], max_workers=8)
            if weight_path.exists():
                break
        except Exception as exc:
            last_error = exc
            print('download failed:', repr(exc))
    if not weight_path.exists():
        raise RuntimeError(f'Cannot download layout weight: {last_error}')
print('weight_path =', weight_path, weight_path.stat().st_size)
"""
    ),
    code(
        r"""import fitz
from tqdm.auto import tqdm


def render_one_page(args):
    pdf_path, page_index, scale, out_dir = args
    doc = fitz.open(str(pdf_path))
    page_no = page_index + 1
    out_path = Path(out_dir) / f'page_{page_no:03d}.jpg'
    if not out_path.exists():
        pix = doc[page_index].get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pix.save(str(out_path))
    return {'page_number': page_no, 'path': str(out_path)}


pdf_doc = fitz.open(str(INPUT_PDF))
total_pages = len(pdf_doc)
start_idx = max(0, START_PAGE - 1)
end_idx = total_pages if END_PAGE is None else min(total_pages, END_PAGE)
page_indexes = list(range(start_idx, end_idx))
print('total_pages =', total_pages, 'selected =', len(page_indexes))

render_jobs = [(INPUT_PDF, i, RENDER_SCALE, PAGES_DIR) for i in page_indexes]
page_images = []
workers = min(CPU_COUNT, max(1, len(render_jobs)))
with ProcessPoolExecutor(max_workers=workers) as executor:
    futures = [executor.submit(render_one_page, job) for job in render_jobs]
    for future in tqdm(as_completed(futures), total=len(futures), desc='Render PDF pages'):
        page_images.append(future.result())
page_images = sorted(page_images, key=lambda item: item['page_number'])
print('rendered pages:', len(page_images))
"""
    ),
    code(
        r"""from PIL import Image, ImageDraw

ID_TO_NAME = {0: 'title', 1: 'plain text', 2: 'abandon', 3: 'figure', 4: 'figure_caption', 5: 'table', 6: 'table_caption', 7: 'table_footnote', 8: 'isolate_formula', 9: 'formula_caption'}
COLORS = {'title': '#2F80ED', 'plain text': '#27AE60', 'abandon': '#828282', 'figure': '#EB5757', 'figure_caption': '#F2994A', 'table': '#9B51E0', 'table_caption': '#BB6BD9', 'table_footnote': '#56CCF2', 'isolate_formula': '#219653', 'formula_caption': '#6FCF97'}


def resolve_device():
    if DEVICE != 'auto':
        return DEVICE
    try:
        import torch
        return '0' if torch.cuda.is_available() else 'cpu'
    except Exception:
        return 'cpu'


def load_layout_model():
    config = {
        'model_path': str(weight_path),
        'img_size': IMG_SIZE,
        'conf_thres': CONF_THRES,
        'iou_thres': IOU_THRES,
        'visualize': False,
        'device': LAYOUT_DEVICE,
    }
    try:
        from pdf_extract_kit.tasks.layout_detection.models.yolo import LayoutDetectionYOLO
        print('Using PDF-Extract-Kit LayoutDetectionYOLO wrapper')
        return 'pdf_extract_kit.LayoutDetectionYOLO', LayoutDetectionYOLO(config)
    except Exception as exc:
        print('PDF-Extract-Kit wrapper failed, fallback to direct YOLO:', repr(exc))

    try:
        from doclayout_yolo import YOLOv10
        print('Using doclayout_yolo.YOLOv10 fallback')
        return 'doclayout_yolo.YOLOv10', YOLOv10(str(weight_path))
    except Exception as exc:
        print('YOLOv10 load failed, fallback to ultralytics.YOLO:', repr(exc))
        from ultralytics import YOLO
        return 'ultralytics.YOLO', YOLO(str(weight_path))


LAYOUT_DEVICE = resolve_device()
layout_engine, layout_model = load_layout_model()
print('LAYOUT_DEVICE =', LAYOUT_DEVICE)
print('layout_engine =', layout_engine)


def predict_one_layout(page_image_path):
    page_image_path = str(page_image_path)
    if layout_engine == 'pdf_extract_kit.LayoutDetectionYOLO':
        return layout_model.predict([page_image_path], str(LAYOUT_VIS_DIR), image_ids=[Path(page_image_path).stem])[0]
    return layout_model.predict(page_image_path, imgsz=IMG_SIZE, conf=CONF_THRES, iou=IOU_THRES, verbose=False, device=LAYOUT_DEVICE)[0]


def result_to_detections(result):
    boxes_obj = result.boxes
    if boxes_obj is None or len(boxes_obj) == 0:
        return []
    xyxy = boxes_obj.xyxy.detach().cpu().numpy()
    cls = boxes_obj.cls.detach().cpu().numpy().astype(int)
    conf = boxes_obj.conf.detach().cpu().numpy()
    dets = []
    for idx, (box, klass, score) in enumerate(zip(xyxy, cls, conf)):
        x0, y0, x1, y1 = [float(v) for v in box]
        label = ID_TO_NAME.get(int(klass), str(int(klass)))
        dets.append({'det_index': idx, 'category_id': int(klass), 'category_type': label, 'score': float(score), 'bbox': [round(x0, 2), round(y0, 2), round(x1, 2), round(y1, 2)], 'poly': [round(x0, 2), round(y0, 2), round(x1, 2), round(y0, 2), round(x1, 2), round(y1, 2), round(x0, 2), round(y1, 2)], 'source': 'pdf_extract_kit_doclayout_yolo'})
    return sorted(dets, key=lambda d: (d['bbox'][1], d['bbox'][0]))


def draw_layout(page_image_path, detections, out_path):
    image = Image.open(page_image_path).convert('RGB')
    draw = ImageDraw.Draw(image)
    for det in detections:
        x0, y0, x1, y1 = det['bbox']
        label = det['category_type']
        color = COLORS.get(label, '#000000')
        draw.rectangle((x0, y0, x1, y1), outline=color, width=3)
        draw.text((x0 + 3, max(0, y0 - 14)), f'{label} {det["score"]:.2f}', fill=color)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(out_path, quality=90)


layout_by_page = {}
for page in tqdm(page_images, desc='PDF-Extract-Kit layout detection'):
    page_no = int(page['page_number'])
    result = predict_one_layout(page['path'])
    detections = result_to_detections(result)
    layout_by_page[page_no] = detections
    draw_layout(page['path'], detections, LAYOUT_VIS_DIR / f'page_{page_no:03d}_layout.jpg')
    counts = {k: sum(1 for d in detections if d['category_type'] == k) for k in sorted(VISUAL_TYPES | CAPTION_TYPES)}
    print('page', page_no, 'detections', len(detections), counts)
"""
    ),
    code(
        r"""def clean_text(value):
    return re.sub(r'\s+', ' ', str(value or '').replace('\u00a0', ' ')).strip()


def bbox_area(bbox):
    x0, y0, x1, y1 = bbox
    return max(0.0, x1 - x0) * max(0.0, y1 - y0)


def bbox_center(bbox):
    x0, y0, x1, y1 = bbox
    return ((x0 + x1) / 2, (y0 + y1) / 2)


def clamp_bbox(bbox, width, height, pad=0):
    x0, y0, x1, y1 = [int(round(v)) for v in bbox]
    return (max(0, x0 - pad), max(0, y0 - pad), min(width, x1 + pad), min(height, y1 + pad))


def nearest_caption(visual, captions, used_caption_indexes, page_width, page_height):
    vx, _ = bbox_center(visual['bbox'])
    vbox = visual['bbox']
    target_type = 'table_caption' if visual['category_type'] == 'table' else 'figure_caption'
    candidates = []
    for idx, cap in enumerate(captions):
        if idx in used_caption_indexes or cap['category_type'] != target_type:
            continue
        cx, _ = bbox_center(cap['bbox'])
        horizontal = abs(cx - vx) / max(1, page_width)
        if cap['bbox'][1] >= vbox[3]:
            vertical_gap = cap['bbox'][1] - vbox[3]
            direction_penalty = 0.0
        else:
            vertical_gap = vbox[1] - cap['bbox'][3]
            direction_penalty = 0.12
        if horizontal > 0.38 or vertical_gap > page_height * 0.22:
            continue
        score = horizontal * 1.6 + vertical_gap / max(1, page_height) + direction_penalty
        candidates.append((score, idx, cap))
    if not candidates:
        return None, None
    _, idx, cap = min(candidates, key=lambda item: item[0])
    used_caption_indexes.add(idx)
    return idx, cap


image_records = []
block_records = []
page_records = []
for page in page_images:
    page_no = int(page['page_number'])
    page_image_path = Path(page['path'])
    image = Image.open(page_image_path).convert('RGB')
    width, height = image.size
    detections = layout_by_page.get(page_no, [])
    min_area = width * height * MIN_CROP_AREA_RATIO
    captions = [d for d in detections if d['category_type'] in CAPTION_TYPES]
    used_caption_indexes = set()
    page_blocks = []
    visual_index = 1

    for order, det in enumerate(detections):
        block = dict(det)
        block.update({'page': page_no, 'page_number': page_no, 'order': order})
        if det['category_type'] in VISUAL_TYPES and bbox_area(det['bbox']) >= min_area:
            _, caption = nearest_caption(det, captions, used_caption_indexes, width, height)
            visual_type = det['category_type']
            stem = f'page_{page_no:03d}_{visual_type}_{visual_index:02d}'
            x0, y0, x1, y1 = clamp_bbox(det['bbox'], width, height, pad=PAD_PX)
            crop_path = IMAGES_DIR / f'{stem}.jpg'
            image.crop((x0, y0, x1, y1)).save(crop_path, quality=92)
            caption_crop_rel = ''
            caption_bbox = None
            if caption and SAVE_CAPTION_CROPS:
                cx0, cy0, cx1, cy1 = clamp_bbox(caption['bbox'], width, height, pad=PAD_PX)
                caption_crop = CAPTIONS_DIR / f'{stem}_caption.jpg'
                image.crop((cx0, cy0, cx1, cy1)).save(caption_crop, quality=92)
                caption_crop_rel = caption_crop.relative_to(OUTPUT_DIR).as_posix()
                caption_bbox = caption['bbox']
            block.update({'type': 'table' if visual_type == 'table' else 'image', 'id': stem, 'label': stem, 'path': crop_path.relative_to(OUTPUT_DIR).as_posix(), 'caption': '', 'caption_bbox': caption_bbox, 'caption_crop_path': caption_crop_rel})
            image_records.append(block)
            visual_index += 1
        else:
            block.update({'type': det['category_type']})
        page_blocks.append(block)
        block_records.append(block)

    try:
        page_text = clean_text(pdf_doc[page_no - 1].get_text('text'))
    except Exception:
        page_text = ''
    page_record = {'source_pdf': str(INPUT_PDF), 'page_number': page_no, 'page_image_path': page_image_path.relative_to(OUTPUT_DIR).as_posix(), 'width': width, 'height': height, 'text': page_text, 'layout_dets': detections, 'text_blocks': page_blocks, 'engine': 'pdf_extract_kit_doclayout_yolo'}
    page_records.append(page_record)
    (METADATA_DIR / 'pages').mkdir(parents=True, exist_ok=True)
    (METADATA_DIR / 'pages' / f'page_{page_no:03d}.json').write_text(json.dumps(page_record, ensure_ascii=False, indent=2), encoding='utf-8')

(METADATA_DIR / 'layout_detections.json').write_text(json.dumps(page_records, ensure_ascii=False, indent=2), encoding='utf-8')
(METADATA_DIR / 'images.json').write_text(json.dumps(image_records, ensure_ascii=False, indent=2), encoding='utf-8')
(METADATA_DIR / 'blocks.jsonl').write_text('\n'.join(json.dumps(x, ensure_ascii=False) for x in block_records) + '\n', encoding='utf-8')
print('pages:', len(page_records), 'images/tables:', len(image_records), 'blocks:', len(block_records))
"""
    ),
    code(
        r"""def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    text = clean_text(text)
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(0, end - overlap)
    return [chunk for chunk in chunks if chunk]


book_lines = [f'# {INPUT_PDF.stem}', '', f'> Source PDF: `{INPUT_PDF}`', f'> Generated at: `{datetime.now(timezone.utc).isoformat()}`', f'> Engine: `PDF-Extract-Kit DocLayout-YOLO layout only`', '> Note: notebook này crop ảnh/bảng/caption; text OCR tiếng Việt cần merge từ DeepSeek-OCR.', '']
rag_records = []
for page in page_records:
    page_no = int(page['page_number'])
    book_lines += [f'## PDF Page {page_no}', '']
    page_images_records = [r for r in image_records if int(r['page_number']) == page_no]
    if page.get('text'):
        book_lines += [str(page['text']), '']
    for record in page_images_records:
        alt = record.get('label') or record.get('id')
        book_lines += [f'![{alt}]({record["path"]})', '']
        if INCLUDE_CAPTION_CROPS_IN_MD and record.get('caption_crop_path'):
            book_lines += [f'![{alt} caption]({record["caption_crop_path"]})', '']
    book_lines += ['</break>', '']
    image_refs = [{k: r.get(k) for k in ['id', 'path', 'caption_crop_path', 'label', 'type']} for r in page_images_records]
    for idx, chunk in enumerate(chunk_text(str(page.get('text') or '')), start=1):
        rag_records.append({'chunk_id': f'page_{page_no:03d}_{idx:02d}', 'source_pdf': str(INPUT_PDF), 'page_number': page_no, 'text': chunk, 'images': image_refs, 'engine': 'pdf_extract_kit_doclayout_yolo'})
    if not page.get('text') and image_refs:
        rag_records.append({'chunk_id': f'page_{page_no:03d}_layout_01', 'source_pdf': str(INPUT_PDF), 'page_number': page_no, 'text': '', 'images': image_refs, 'engine': 'pdf_extract_kit_doclayout_yolo'})

(OUTPUT_DIR / 'book.md').write_text('\n'.join(book_lines).strip() + '\n', encoding='utf-8')
(OUTPUT_DIR / 'rag_chunks.jsonl').write_text('\n'.join(json.dumps(x, ensure_ascii=False) for x in rag_records) + ('\n' if rag_records else ''), encoding='utf-8')
summary = {'source_pdf': str(INPUT_PDF), 'generated_at': datetime.now(timezone.utc).isoformat(), 'page_range': {'start': START_PAGE, 'end': END_PAGE}, 'pages_processed': len(page_records), 'engine': {'layout': 'pdf_extract_kit_doclayout_yolo', 'ocr': 'none'}, 'model': {'weight': str(weight_path), 'img_size': IMG_SIZE, 'conf_thres': CONF_THRES, 'iou_thres': IOU_THRES, 'device': LAYOUT_DEVICE}, 'stats': {'blocks': len(block_records), 'images': len([r for r in image_records if r.get('type') == 'image']), 'tables': len([r for r in image_records if r.get('type') == 'table']), 'rag_chunks': len(rag_records)}, 'outputs': {'markdown': 'book.md', 'rag_chunks': 'rag_chunks.jsonl', 'page_metadata_dir': 'metadata/pages', 'block_metadata': 'metadata/blocks.jsonl', 'image_metadata': 'metadata/images.json', 'layout_metadata': 'metadata/layout_detections.json'}}
(METADATA_DIR / 'book.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(summary, ensure_ascii=False, indent=2))
"""
    ),
    code(
        r"""from docx import Document
from docx.shared import Inches


def make_docx():
    doc = Document()
    doc.add_heading(INPUT_PDF.stem, level=1)
    for line in (OUTPUT_DIR / 'book.md').read_text(encoding='utf-8').splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith('> ') or stripped == '</break>':
            continue
        if stripped.startswith('# '):
            continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith('!['):
            match = re.search(r'\]\((.*?)\)', stripped)
            if match:
                image_path = OUTPUT_DIR / match.group(1)
                if image_path.exists():
                    try:
                        doc.add_picture(str(image_path), width=Inches(5.6))
                    except Exception:
                        doc.add_paragraph(str(image_path))
        else:
            doc.add_paragraph(stripped)
    doc.save(str(OUTPUT_DIR / 'book.docx'))
    return True


def make_contact_sheet(records):
    records = [r for r in records if r.get('path')]
    if not records:
        return None
    thumb_w, label_h, pad, cols = 260, 42, 12, 3
    rows = math.ceil(len(records) / cols)
    sheet = Image.new('RGB', (cols * thumb_w + (cols + 1) * pad, rows * (thumb_w + label_h) + (rows + 1) * pad), 'white')
    draw = ImageDraw.Draw(sheet)
    for idx, rec in enumerate(records):
        path = OUTPUT_DIR / rec['path']
        if not path.exists():
            continue
        with Image.open(path).convert('RGB') as img:
            img.thumbnail((thumb_w, thumb_w - label_h), Image.Resampling.LANCZOS)
            col, row = idx % cols, idx // cols
            x = pad + col * (thumb_w + pad)
            y = pad + row * (thumb_w + label_h + pad)
            sheet.paste(img, (x + (thumb_w - img.width) // 2, y))
            draw.text((x + 4, y + thumb_w - 12), f'{rec.get("id")} p.{rec.get("page_number")}', fill=(20, 20, 20))
    out = PREVIEW_DIR / 'images_contact.jpg'
    sheet.save(out, quality=88)
    return out


print('docx:', make_docx())
contact = make_contact_sheet(image_records)
print('contact:', contact)
for path in [OUTPUT_DIR / 'book.md', OUTPUT_DIR / 'book.docx', OUTPUT_DIR / 'rag_chunks.jsonl', METADATA_DIR / 'book.json', METADATA_DIR / 'images.json', contact]:
    if path and Path(path).exists():
        print(path, Path(path).stat().st_size)
"""
    ),
    code(
        r"""zip_path = OUTPUT_DIR.with_suffix('.zip')
if ZIP_OUTPUT:
    with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in sorted(OUTPUT_DIR.rglob('*')):
            if path.is_file():
                archive.write(path, path.relative_to(OUTPUT_DIR.parent))
    print('zip:', zip_path, zip_path.stat().st_size)

try:
    from IPython.display import display, Image as IPImage
    contact = PREVIEW_DIR / 'images_contact.jpg'
    layout_files = sorted(LAYOUT_VIS_DIR.glob('*_layout.jpg'))
    if contact.exists():
        display(IPImage(filename=str(contact)))
    if layout_files:
        display(IPImage(filename=str(layout_files[0])))
except Exception as exc:
    print(exc)

print('\nImportant files:')
for rel in ['book.md', 'book.docx', 'rag_chunks.jsonl', 'metadata/book.json', 'metadata/images.json', 'metadata/layout_detections.json', '_previews/images_contact.jpg']:
    path = OUTPUT_DIR / rel
    print(rel, 'OK' if path.exists() else 'MISSING', path.stat().st_size if path.exists() else '')

print('\nFinal output directory:')
print(OUTPUT_DIR)
"""
    ),
]

nb = {
    "cells": cells,
    "metadata": {
        "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
        "language_info": {"name": "python", "pygments_lexer": "ipython3"},
    },
    "nbformat": 4,
    "nbformat_minor": 5,
}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(nb, ensure_ascii=False, indent=1), encoding="utf-8")
print(OUT)
